"""Generate doc/AbCS_Fall_2026_Tester_Review.xlsx for tester feedback."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "doc" / "AbCS_Fall_2026_Tester_Review.xlsx"

# Item ID, category, enhancement, description, effort, risk, risk reason, timing
# Excel order: description, then tester columns, then effort/risk/timing (developer notes).
ROWS = [
    (
        "C01",
        "Core - Wave 0",
        "Schema batch (foundation)",
        "Behind-the-scenes database upgrade that adds storage for Want to Read, "
        "ratings, cover images, and collection folder paths. You will not see "
        "new buttons yet, but import, backup, and older databases should still work.",
        "2-3 days",
        "Medium",
        "Touches the database all features use; must not break import or legacy backup/restore.",
        "First - prerequisite for Waves 1-3",
    ),
    (
        "C02",
        "Core - Wave 1",
        "Want to Read",
        "Mark books you plan to listen to without moving them to another collection. "
        "Filter the main list to show only your to-be-read queue. The flag clears when "
        "you mark a book as read.",
        "3.5-4 days",
        "Low",
        "Simple flag and filter; main risk is shortcut and filter accessibility.",
        "After Wave 0 schema",
    ),
    (
        "C03",
        "Core - Wave 1",
        "Open audiobook location",
        "From Book Details, open the folder where the audiobook files live (or show "
        "the file in your file manager) so you can play them in your own player. "
        "Not an in-app player.",
        "~1 day",
        "Low",
        "Small OS-specific open-folder behavior; button disabled when path is missing.",
        "Ships with Want to Read",
    ),
    (
        "C04",
        "Core - Wave 2",
        "Book ratings",
        "Store a numeric rating (for example 4.5) on each book. Type your own rating "
        "or fill it from web metadata. See ratings in the main book list and Book Details.",
        "4-5 days",
        "Medium",
        "One-time migration from old Rating text in plot comments; web metadata save must not regress.",
        "After Wave 0; combine sprint with covers",
    ),
    (
        "C05",
        "Core - Wave 2",
        "Cover images",
        "When you fetch web metadata for a book, save its cover image and show it in "
        "Book Details and Import Detail. Covers are not shown in the main book table.",
        "Part of 5-6 days with backups",
        "Medium",
        "Network download on save; cover files on disk; deleting a book must remove its cover file.",
        "After Wave 0; same sprint as ratings",
    ),
    (
        "C06",
        "Core - Wave 2",
        "Better backups (zip package)",
        "New backups become a zip file containing the database plus cover images. "
        "Restore brings covers back with your library. Old database-only backups still "
        "restore without covers.",
        "Part of 5-6 days with covers",
        "Medium",
        "Must preserve legacy database restore; full reset must clean the covers folder.",
        "Ships with cover images",
    ),
    (
        "C07",
        "Core - Wave 3",
        "Collection library folder",
        "Optional folder path for each collection (for example where your Audible rips "
        "live). Import and rescan default to that folder. No files are moved.",
        "2-3 days",
        "Low",
        "Optional path on collections; import pre-fill only.",
        "After Wave 0; before rescan",
    ),
    (
        "C08",
        "Core - Wave 3",
        "Rescan / update from folder",
        "Scan a folder again and update books already in your library (length, file path, "
        "track count, and similar) instead of only adding new ones. You choose what to "
        "update. Tag overwrite is off by default.",
        "8-10 days",
        "Medium",
        "Wrong matches could update the wrong book; ambiguous matches are flagged; tag overwrite is opt-in.",
        "After collection library folder",
    ),
    (
        "C09",
        "Core - Wave 4 optional",
        "Name consistency check",
        "Find similar spellings of author names, titles, and genres (for example Connolly "
        "vs Connelly) and help you merge them into one spelling. Similar to Duplicate Check. "
        "No automatic silent merges.",
        "2-3 weeks",
        "Medium",
        "Merge operations can change many books at once; user must confirm each group.",
        "After Wave 3 rescan",
    ),
    (
        "C10",
        "Core - Wave 5",
        "Multiple languages",
        "UI text, menus, and messages in languages such as French and Spanish, with English "
        "as the default. Help docs would be translated separately.",
        "3-5 weeks",
        "Medium",
        "Large mechanical change across about 1,500 strings; accessibility stays safe if names stay paired with labels.",
        "After English feature freeze; separate release",
    ),
    (
        "C11",
        "Deferred / later",
        "Organize files into library folder",
        "Optional wizard to copy or move audiobook folders into a tidy layout under your "
        "collection folder and update paths in AbCS. Preview and dry-run before any file changes.",
        "10-12 days",
        "High",
        "Moves or copies real files on disk; partial failure and rollback are complex. Deferred from fall core.",
        "After rescan A+B are stable; post-fall if needed",
    ),
    (
        "F01",
        "Follow-on",
        "Path health report",
        "A report listing books whose file path no longer exists on disk (moved, deleted, or wrong drive).",
        "2-3 days",
        "Low",
        "Read-only scan and report; no automatic repair in v1.",
        "Wave 4; pairs with export library metadata",
    ),
    (
        "F02",
        "Follow-on",
        "Export library to spreadsheet",
        "Export your book list to CSV or similar for Excel, backup, or sharing - the reverse of Import Book List.",
        "~2 days",
        "Low",
        "Read-only export; main risk is large-library performance.",
        "Wave 4; pairs with path health report",
    ),
    (
        "F03",
        "Follow-on",
        "Missing info filters",
        "Quick filters on the main window: show only books with no plot, no cover, no rating, or no file path.",
        "2-3 days",
        "Medium",
        "Cover and rating filters depend on Wave 2 schema; overlaps conceptually with path health.",
        "After Wave 2",
    ),
    (
        "F04",
        "Follow-on",
        "Mark several books Want to Read",
        "Select multiple books on the main list and mark or clear Want to Read in one step.",
        "~1 day",
        "Low",
        "Simple bulk update using the same field as Book Details.",
        "After Wave 1",
    ),
    (
        "F05",
        "Follow-on",
        "Want to Read during import",
        "Set Want to Read while reviewing a book in Import Detail, before it is added to the library.",
        "~1 day",
        "Low",
        "Extends existing import detail flow with one checkbox.",
        "After Wave 1",
    ),
    (
        "F06",
        "Follow-on",
        "More bulk update options",
        "Extend the Update window so you can change Want to Read, reader, or year for many selected books at once.",
        "2-3 days",
        "Medium",
        "Tri-state bulk edit across multiple fields increases regression surface.",
        "After Wave 1",
    ),
    (
        "F07",
        "Follow-on",
        "Backup reminder",
        "A gentle reminder to create a backup if you have not done so in a while. You choose "
        "whether to act. Prompt only - no silent auto-backup.",
        "1-2 days",
        "Low",
        "User-controlled dismiss and snooze; uses existing backup list.",
        "After Wave 2 zip backup",
    ),
    (
        "F08",
        "Follow-on",
        "Richer statistics",
        "Statistics screen counts for Want to Read, average rating, books with covers, and similar.",
        "1-2 days",
        "Low",
        "Display-only query extensions.",
        "After Wave 2",
    ),
    (
        "F09",
        "Follow-on",
        "Filter by narrator",
        "Show only books read by a chosen narrator or reader.",
        "~2 days",
        "Low",
        "Standard filter pattern on an existing reader field.",
        "Wave 4; no schema dependency",
    ),
    (
        "F10",
        "Follow-on",
        "Series book number",
        "Store book 3 in the series as its own field instead of only in the title.",
        "2-3 days",
        "Medium",
        "Schema migration and careful title/number split from existing data.",
        "Wave 0 add-on or Wave 4",
    ),
    (
        "F11",
        "Follow-on",
        "Export / import settings",
        "Save your preferences to a file and load them on another computer.",
        "~2 days",
        "Medium",
        "Import overwrite needs clear confirmation; key list must stay in sync.",
        "Anytime",
    ),
    (
        "F12",
        "Follow-on - Post-fall",
        "Fetch web info for many books",
        "Queue web metadata fetch for a selection of books with progress and cancel - instead "
        "of one book at a time. Per-book review remains recommended.",
        "1-2 weeks",
        "High",
        "Network-heavy; API rate limits; long runs on large selections.",
        "Post-fall; benefits from background-thread work first",
    ),
    (
        "F13",
        "Follow-on",
        "Web fetch background thread",
        "Run web metadata fetch on a background thread so the app stays responsive during "
        "network calls. Cancel still works. Progress announcements stay on the GUI thread.",
        "3-5 days",
        "High",
        "First production background thread in AbCS; thread safety and cancel must be careful.",
        "After web fetch Phases 1-5 already shipped in 2.10",
    ),
    (
        "F14",
        "Follow-on - Post-fall",
        "Better plot search",
        "Faster search inside long plot summaries on very large libraries.",
        "3-5 days",
        "Medium",
        "Full-text index sync and rebuild add complexity.",
        "Post-fall; most useful on very large libraries",
    ),
    (
        "F15",
        "Follow-on",
        "Import window action toolbar",
        "Import window gets a labeled action toolbar like the main window for common import "
        "actions. Existing shortcuts and enable/disable rules must still work.",
        "~1 day",
        "Low",
        "UI layout polish; reuse main toolbar pattern; no new business logic.",
        "Anytime; deferred from June visual appeal",
    ),
    (
        "F16",
        "Follow-on",
        "Preferences mini-toolbar",
        "Preferences gets a small toolbar for Save, Restore defaults, and Close. Existing Alt "
        "shortcuts and unsaved-changes prompts must remain correct.",
        "0.5-1 day",
        "Low",
        "Visual/layout change only; existing handlers reused.",
        "Anytime; deferred from June visual appeal",
    ),
    (
        "B01",
        "Backlog",
        "Import from other apps",
        "Import a library export from another tool (for example Libib or similar CSV formats) "
        "with less manual editing.",
        "1-2 weeks per format",
        "High",
        "Third-party export formats change; ongoing maintenance per adapter.",
        "Demand-driven after core waves",
    ),
    (
        "B02",
        "Backlog",
        "Saved smart lists",
        "Save a combination of filters (for example unread sci-fi with a plot) and reuse it "
        "with one click.",
        "1-2 weeks",
        "Medium",
        "Users may confuse saved smart lists with normal collections.",
        "After core filter features land",
    ),
    (
        "B03",
        "Backlog",
        "Reading progress",
        "Remember how far you got in a book (not just finished or not). Mainly useful if deeper "
        "playback support is added later. Not an in-app player.",
        "2-3 weeks+",
        "High",
        "Conflicts with AbCS focus as a collection manager rather than a player.",
        "Demand-driven; deferred unless strong tester demand",
    ),
    (
        "B04",
        "Backlog",
        "Tags on books",
        "Multiple labels per book (for example gift or book club) without changing its collection.",
        "2-3 weeks",
        "Medium",
        "Overlaps conceptually with Want to Read and collections; needs clear guidance.",
        "Demand-driven backlog",
    ),
    (
        "B05",
        "Backlog",
        "Check for updates",
        "Help menu option to see if a newer AbCS version is available and open the download page. "
        "No silent install.",
        "2-3 days",
        "Low",
        "Offline users see nothing; no telemetry in v1.",
        "Anytime; good small release win",
    ),
    (
        "M01",
        "Maintenance - not user-facing",
        "CI / test hardening",
        "Strengthen automated regression testing for fall waves. Supports safer development. "
        "Not a user-visible feature.",
        "1-2 days setup + ongoing",
        "Low",
        "Infrastructure work; does not change app features for end users.",
        "Between waves; early in fall cycle",
    ),
    (
        "M02",
        "Maintenance - not user-facing",
        "Dead-code cleanup",
        "Periodic cleanup of unused code to keep the project easier to maintain. "
        "Not a user-visible feature.",
        "0.5-1 day per pass",
        "Low",
        "Only removes confirmed unused code after tests pass.",
        "Between waves or before release tags",
    ),
]

HEADERS = [
    "Item ID",
    "Category / Wave",
    "Enhancement",
    "Brief description",
    "Should we do it? (type X)",
    "Tester priority",
    "Tester comments",
    "Approximate effort",
    "Risk",
    "Risk reason",
    "Timing / dependency",
]

INSTRUCTIONS = [
    "AbCS Fall 2026 Tester Review",
    "",
    "Purpose",
    "This workbook lists planned AbCS improvements for fall 2026 review. "
    "Nothing here is a commitment to build or a fixed release date.",
    "",
    "How to fill Fall Plans",
    "1. Read each Brief description.",
    "2. In Should we do it?, type X if you think AbCS should build it. Leave blank if unsure or no.",
    "3. In Tester priority, type a number for your top choices (1 = most important). "
    "Leave blank if not ranking that item.",
    "4. Use Tester comments for notes, accessibility concerns, or why you would or would not use it.",
    "5. Skip Approximate effort, Risk, Risk reason, and Timing / dependency. Those are developer notes.",
    "",
    "Column meanings",
    "The three tester columns sit right after Brief description so you can fill them without scrolling past developer notes.",
    "Approximate effort: developer estimate only.",
    "Risk: implementation or data risk (Low, Medium, or High), not danger to you as a user.",
    "Timing/dependency: when the work fits relative to other items.",
    "Maintenance rows are marked not user-facing. Skip them if you only want to rank features.",
    "",
    "Feedback tips",
    "Tell us your top 3 to 5 items overall.",
    "Tell us anything here you would not use.",
    "Tell us anything missing from this list that you would use regularly.",
    "",
    "Accessibility note",
    "This sheet uses plain cells (type X) instead of form-control checkboxes so JAWS and NVDA "
    "can fill it reliably.",
    "Header row is frozen. Filters are available on Fall Plans.",
    "",
    "Related docs",
    "doc/abcs_proposed_enhancements.md - plain-language summary",
    "doc/plan_enhancements_fall2026.md - internal schedule",
]


def main() -> None:
    assert len(ROWS) == 34, f"Expected 34 rows, got {len(ROWS)}"

    wb = Workbook()

    ws_i = wb.active
    ws_i.title = "Instructions"
    for i, line in enumerate(INSTRUCTIONS, start=1):
        cell = ws_i.cell(row=i, column=1, value=line)
        if i == 1:
            cell.font = Font(bold=True, size=14)
        elif line in {
            "Purpose",
            "How to fill Fall Plans",
            "Column meanings",
            "Feedback tips",
            "Accessibility note",
            "Related docs",
        }:
            cell.font = Font(bold=True, size=12)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws_i.column_dimensions["A"].width = 110

    ws = wb.create_sheet("Fall Plans")
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row in ROWS:
        ws.append([*row[:4], "", "", "", *row[4:]])

    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    for excel_row in ws.iter_rows(
        min_row=1, max_row=1 + len(ROWS), min_col=1, max_col=len(HEADERS)
    ):
        for cell in excel_row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = thin

    widths = {
        "A": 10,
        "B": 28,
        "C": 34,
        "D": 55,
        "E": 18,
        "F": 14,
        "G": 30,
        "H": 18,
        "I": 10,
        "J": 40,
        "K": 36,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:K{1 + len(ROWS)}"
    ws.row_dimensions[1].height = 30
    for i in range(2, 2 + len(ROWS)):
        ws.row_dimensions[i].height = 60

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Rows: {len(ROWS)}")


if __name__ == "__main__":
    main()
