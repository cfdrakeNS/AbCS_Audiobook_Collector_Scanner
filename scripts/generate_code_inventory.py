"""Generate a dated Excel inventory of AbCS Python modules.

Scans src/, test/, scripts/, and root utility .py files. Line counts include
blank lines and comments. Modified dates come from the local filesystem.

Output: doc/abcs_code_inventory_MonDD_YYYY.xlsx

Run:
    python scripts/generate_code_inventory.py
"""

from __future__ import annotations

import ast
import datetime as dt
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DOC = ROOT / "doc"
SKIP_DIRS = {
    ".git",
    ".venv",
    "venv",
    "__pycache__",
    "archive",
    "build",
    "dist",
    "data",
    "graphics",
    "Graphics",
    "help_docs",
    "doc",
    "linux",
}

ARCHIVE = ROOT / "archive"
OLD_INVENTORY_CANDIDATES = [
    ARCHIVE / "abcs_code_inventory_June24_2026.md",
    DOC / "abcs_code_inventory_June24_2026.md",
]
TABLE_ROW_RE = re.compile(
    r"^\| `([^`]+)` \| (.+?) \| [0-9,]+ \| \d{4}-\d{2}-\d{2} \|$"
)
SECTION_RE = re.compile(r"^#{2,3} (.+)$")

NOT_INCLUDED = [
    ("archive/", "Archived demos, old tests, and retired scripts — not active source"),
    (
        "External: pyside6-accessible-ui-reference",
        "Standalone PySide6 accessibility reference (separate GitHub repo)",
    ),
    ("build/, dist/", "PyInstaller build output"),
    ("data/", "SQLite database and backups (runtime data)"),
    ("graphics/", "Icons, splash, and image assets"),
    ("test/fixtures/", "SQL schema fixtures for tests"),
    ("build_installer.bat, build_installer.iss", "Windows installer build"),
    ("AbCS.spec, requirements.txt", "Packaging and dependency config"),
    ("build_linux.sh, build_linux_debug.sh, build_linux_common.sh, build_linux_legacy_hp.sh", "Linux PyInstaller wrappers"),
    ("scripts/*.ps1", "GitHub release, history, and repo-cleanup PowerShell helpers"),
    ("Documentation (doc/, help_docs/, README.md, etc.)", "Markdown guides, not code modules"),
]


def line_count(path: Path) -> int:
    return len(path.read_text(encoding="utf-8", errors="replace").splitlines())


def modified_date(path: Path) -> str:
    return dt.datetime.fromtimestamp(path.stat().st_mtime).date().isoformat()


def first_docstring_line(path: Path) -> str:
    try:
        tree = ast.parse(path.read_text(encoding="utf-8"))
    except SyntaxError:
        return ""
    doc = ast.get_docstring(tree)
    if not doc:
        return ""
    for raw in doc.strip().splitlines():
        line = raw.strip()
        if line:
            return line
    return ""


def load_old_descriptions() -> dict[str, str]:
    """Map posix relative path -> description from the June 24 markdown inventory."""
    inventory = next((p for p in OLD_INVENTORY_CANDIDATES if p.is_file()), None)
    if inventory is None:
        return {}
    folder = ""
    out: dict[str, str] = {}
    for line in inventory.read_text(encoding="utf-8").splitlines():
        sec = SECTION_RE.match(line)
        if sec:
            title = sec.group(1)
            if title.startswith("`test/`") or title.startswith("test/"):
                folder = "test/"
            elif title.startswith("`scripts/`") or title.startswith("scripts/"):
                folder = "scripts/"
            elif "Project root" in title:
                folder = ""
            elif title.startswith("src/") or title.startswith("`src/"):
                cleaned = title.strip("`")
                cleaned = cleaned.split("(")[0].strip()
                if "root" in cleaned:
                    folder = "src/"
                else:
                    folder = cleaned if cleaned.endswith("/") else f"{cleaned}/"
            continue
        row = TABLE_ROW_RE.match(line)
        if not row:
            continue
        name, desc = row.group(1), row.group(2).strip()
        rel = f"{folder}{name}" if folder else name
        out[rel] = desc
        out[name] = desc
    return out


def classify(rel: Path) -> tuple[str, str]:
    parts = rel.parts
    if parts[0] == "src":
        if len(parts) == 2:
            return "src", "src/"
        return "src", f"src/{parts[1]}/"
    if parts[0] == "test":
        return "test", "test/"
    if parts[0] == "scripts":
        return "scripts", "scripts/"
    return "root", "(project root)"


def collect_python_files() -> list[Path]:
    files: list[Path] = []
    for path in ROOT.rglob("*.py"):
        rel = path.relative_to(ROOT)
        if any(part in SKIP_DIRS for part in rel.parts):
            continue
        area, _folder = classify(rel)
        if area == "root" and rel.parts[0] not in {"get_version.py"}:
            continue
        files.append(path)
    files.sort(key=lambda p: p.relative_to(ROOT).as_posix().lower())
    return files


def describe(rel_posix: str, name: str, path: Path, old: dict[str, str]) -> str:
    doc = first_docstring_line(path)
    if doc:
        return doc
    return old.get(rel_posix) or old.get(name) or "Python module"


def inventory_rows() -> list[tuple[str, str, str, str, int, str]]:
    old = load_old_descriptions()
    rows = []
    for path in collect_python_files():
        rel = path.relative_to(ROOT)
        area, folder = classify(rel)
        rel_posix = rel.as_posix()
        rows.append(
            (
                area,
                folder,
                path.name,
                describe(rel_posix, path.name, path, old),
                line_count(path),
                modified_date(path),
            )
        )
    return rows


def apply_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def style_table(ws, row_count: int, col_count: int, data_height: int = 32) -> None:
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    for excel_row in ws.iter_rows(min_row=1, max_row=row_count, min_col=1, max_col=col_count):
        for cell in excel_row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = thin
    ws.row_dimensions[1].height = 22
    for i in range(2, row_count + 1):
        ws.row_dimensions[i].height = data_height
    ws.freeze_panes = "A2"
    last_col = get_column_letter(col_count)
    ws.auto_filter.ref = f"A1:{last_col}{row_count}"


def set_widths(ws, widths: dict[str, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def write_intro(ws, title: str, lines: list[str]) -> None:
    ws.title = "Summary"
    cell = ws.cell(row=1, column=1, value=title)
    cell.font = Font(bold=True, size=14)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, line in enumerate(lines, start=2):
        c = ws.cell(row=i, column=1, value=line)
        if line in {
            "Purpose",
            "How to read",
            "Counts",
            "Related",
        }:
            c.font = Font(bold=True, size=12)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 110


def main() -> None:
    today = dt.date.today()
    stamp = today.strftime("%b%d_%Y")
    out = DOC / f"abcs_code_inventory_{stamp}.xlsx"

    rows = inventory_rows()
    assert rows, "No Python modules found"

    by_area: dict[str, list[tuple]] = {}
    for row in rows:
        by_area.setdefault(row[0], []).append(row)

    src_rows = by_area.get("src", [])
    src_lines = sum(r[4] for r in src_rows)
    total_lines = sum(r[4] for r in rows)

    wb = Workbook()
    ws_s = wb.active
    write_intro(
        ws_s,
        f"AbCS Code Inventory — {today.strftime('%B %d, %Y')}",
        [
            "",
            "Purpose",
            "Inventory of Python source modules in AbCS (Audiobook Collector Scanner). "
            "Line counts include blank lines and comments. Last modified dates are from "
            "the local filesystem.",
            "",
            "How to read",
            "Use the Modules sheet for the full list. Filters are on. "
            "Largest lists the ten biggest src/ files. Not included lists project files "
            "that are not Python application modules.",
            "This workbook does not include a comparison with the previous inventory.",
            "",
            "Counts",
            f"src/ (application): {len(src_rows)} modules, {src_lines:,} lines",
            f"test/: {len(by_area.get('test', []))} modules, "
            f"{sum(r[4] for r in by_area.get('test', [])):,} lines",
            f"scripts/: {len(by_area.get('scripts', []))} modules, "
            f"{sum(r[4] for r in by_area.get('scripts', [])):,} lines",
            f"Project root utilities: {len(by_area.get('root', []))} modules, "
            f"{sum(r[4] for r in by_area.get('root', [])):,} lines",
            f"Total: {len(rows)} modules, {total_lines:,} lines",
            "",
            "Related",
            "Previous markdown inventory (local archive): archive/abcs_code_inventory_June24_2026.md",
            "Regenerate: python scripts/generate_code_inventory.py",
        ],
    )

    ws_c = wb.create_sheet("Counts")
    apply_header(ws_c, ["Area", "Modules", "Lines"])
    count_data = [
        ("src/", len(src_rows), src_lines),
        (
            "test/",
            len(by_area.get("test", [])),
            sum(r[4] for r in by_area.get("test", [])),
        ),
        (
            "scripts/",
            len(by_area.get("scripts", [])),
            sum(r[4] for r in by_area.get("scripts", [])),
        ),
        (
            "Project root utilities",
            len(by_area.get("root", [])),
            sum(r[4] for r in by_area.get("root", [])),
        ),
        ("Total", len(rows), total_lines),
    ]
    for area, n, lines in count_data:
        ws_c.append([area, n, lines])
        if area == "Total":
            for cell in ws_c[ws_c.max_row]:
                cell.font = Font(bold=True)
    style_table(ws_c, 1 + len(count_data), 3, data_height=22)
    set_widths(ws_c, {"A": 28, "B": 12, "C": 12})

    ws_m = wb.create_sheet("Modules")
    apply_header(
        ws_m,
        ["Area", "Folder", "Module", "Description", "Lines", "Modified"],
    )
    for row in rows:
        ws_m.append(list(row))
    style_table(ws_m, 1 + len(rows), 6, data_height=36)
    set_widths(
        ws_m,
        {"A": 10, "B": 22, "C": 42, "D": 72, "E": 10, "F": 12},
    )

    largest = sorted(src_rows, key=lambda r: r[4], reverse=True)[:10]
    largest_lines = sum(r[4] for r in largest)
    pct = round(100 * largest_lines / src_lines) if src_lines else 0

    ws_l = wb.create_sheet("Largest")
    apply_header(ws_l, ["Rank", "Module", "Lines", "Folder"])
    for i, row in enumerate(largest, start=1):
        _area, folder, name, _desc, lines, _mod = row
        ws_l.append([i, name, lines, folder])
    style_table(ws_l, 1 + len(largest), 4, data_height=22)
    set_widths(ws_l, {"A": 8, "B": 36, "C": 10, "D": 22})
    note_row = 2 + len(largest)
    ws_l.cell(
        row=note_row,
        column=1,
        value=(
            f"These {len(largest)} files account for about {pct}% of all src/ Python lines "
            f"({largest_lines:,} of {src_lines:,})."
        ),
    )
    ws_l.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=4)
    ws_l.cell(row=note_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws_l.row_dimensions[note_row].height = 32

    ws_n = wb.create_sheet("Not included")
    apply_header(ws_n, ["Item", "Notes"])
    for item, notes in NOT_INCLUDED:
        ws_n.append([item, notes])
    style_table(ws_n, 1 + len(NOT_INCLUDED), 2, data_height=28)
    set_widths(ws_n, {"A": 70, "B": 70})

    DOC.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Wrote {out}")
    print(f"Modules: {len(rows)}")
    print(f"Lines: {total_lines:,}")


if __name__ == "__main__":
    main()
